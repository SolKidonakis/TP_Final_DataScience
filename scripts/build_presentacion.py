"""Genera la Presentación Final (~12 slides) en PPTX y PDF."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "reports" / "results.json"
FIGURES = ROOT / "reports" / "figures"
OUT_PPTX = ROOT / "reports" / "Presentacion_Final_Grupo33.pptx"
OUT_PDF = ROOT / "reports" / "Presentacion_Final_Grupo33.pdf"

# Colores
NAVY = RGBColor(0x1F, 0x4E, 0x79)
ACCENT = RGBColor(0xE4, 0x57, 0x56)
DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def pct(x, d=1):
    return f"{100 * float(x):.{d}f}%"


def num(x, d=3):
    return f"{float(x):.{d}f}"


def set_run(run, size=18, bold=False, color=DARK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_title(shape, text, size=28, color=NAVY):
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=True, color=color)


def add_bullets(shape, items, size=16):
    tf = shape.text_frame
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = item
        set_run(run, size=size, color=DARK)


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank


def banner(slide, title):
    # Barra superior
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.9))
    shape.fill.solid()
    shape.fill.fore_color.rgb = NAVY
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "  " + title
    set_run(run, size=22, bold=True, color=WHITE)


def footer(slide, page, total=12):
    box = slide.shapes.add_textbox(Inches(0.4), Inches(7.1), Inches(12.5), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"Grupo 33 · Detección de Fraude · {page}/{total}"
    set_run(run, size=10, color=GRAY)


def try_picture(slide, name, left, top, width):
    path = FIGURES / name
    if path.exists():
        slide.shapes.add_picture(str(path), left, top, width=width)
        return True
    return False


def build_pptx(data):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    eda = data["eda"]
    rand = data["random_split"]
    temp = data["temporal_split"]
    best_name = rand["best_model"]
    best = rand["models"][best_name]
    m50 = best["metrics_thr_0.5"]
    msel = best["metrics_selected"]
    cap = {r["k"]: r for r in best["captura_at_k"]}
    best_t = temp["best_model"]
    m50_t = temp["models"][best_t]["metrics_thr_0.5"]

    # 1 Portada
    s = blank_slide(prs)
    shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = NAVY
    shape.line.fill.background()
    box = s.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(4))
    tf = box.text_frame
    for i, (text, size, bold) in enumerate([
        ("PROYECTO FINAL INTEGRADOR", 16, False),
        ("Detección de Fraude en Tarjetas de Crédito", 32, True),
        ("Grupo 33", 20, True),
        ("Mariana Chaile · Eugenia Salcedo Aisa · Sol Kidonakis", 14, False),
        ("Marco Garcia · Esteban Derghazarian", 14, False),
        ("Octubre 2026", 14, False),
    ]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = text
        set_run(run, size=size, bold=bold, color=WHITE)

    # 2 Agenda
    s = blank_slide(prs)
    banner(s, "Agenda")
    box = s.shapes.add_textbox(Inches(1), Inches(1.4), Inches(11), Inches(5))
    add_bullets(box, [
        "1. Problema de negocio y objetivo",
        "2. Datos y desbalance extremo",
        "3. Metodología CRISP-DM",
        "4. Modelos y métricas principales (PR-AUC, Recall, Precision)",
        "5. Métrica de negocio: captura Top-100 / Top-500",
        "6. Validación temporal vs aleatoria",
        "7. Impacto, riesgos y recomendaciones",
    ], size=18)
    footer(s, 2)

    # 3 Problema
    s = blank_slide(prs)
    banner(s, "El problema: fraude con pocos casos y muchas operaciones")
    box = s.shapes.add_textbox(Inches(0.7), Inches(1.3), Inches(12), Inches(5.5))
    add_bullets(box, [
        "Pérdidas globales > USD 33.000 M / año por fraude en tarjetas.",
        f"Dataset ULB: {eda['n_rows']:,} transacciones · solo {eda['n_fraud']} fraudes ({pct(eda['fraud_rate'], 3)}).",
        "Audiencia: Prevención de Fraude / Riesgo Crediticio.",
        "Objetivo: maximizar detección SIN saturar al equipo con falsas alertas.",
        "Hipótesis: ML supera reglas estáticas priorizando una cola de revisión accionable.",
    ], size=17)
    footer(s, 3)

    # 4 Datos
    s = blank_slide(prs)
    banner(s, "Datos: desbalance extremo")
    try_picture(s, "01_class_distribution.png", Inches(0.6), Inches(1.2), Inches(6.2))
    box = s.shapes.add_textbox(Inches(7.2), Inches(1.5), Inches(5.5), Inches(5))
    add_bullets(box, [
        f"{eda['n_fraud']} fraudes vs {eda['n_rows'] - eda['n_fraud']:,} legítimas.",
        "Sin nulos; duplicados eliminados.",
        "V1–V28 = PCA anonimizado.",
        "Amount sesgado → log1p.",
        "Time habilita validación temporal.",
        "Accuracy engaña → usamos PR-AUC.",
    ], size=16)
    footer(s, 4)

    # 5 Metodología
    s = blank_slide(prs)
    banner(s, "Metodología CRISP-DM")
    box = s.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.5), Inches(5.5))
    add_bullets(box, [
        "Preparación en Pipeline: StandardScaler(Time, Amount_log) sin leakage.",
        "Modelos: Logistic Regression · Random Forest · XGBoost · LightGBM · XGBoost+SMOTE.",
        "Métricas principales: PR-AUC, Recall, Precision (+ F1, ROC-AUC).",
        "Negocio: captura Top-K y volumen de alertas por umbral.",
        "Validación: holdout estratificado 80/20  vs  split temporal por Time.",
    ], size=17)
    footer(s, 5)

    # 6 Resultados clave
    s = blank_slide(prs)
    banner(s, f"Resultados clave — modelo: {best_name}")
    # KPI cards
    kpis = [
        ("PR-AUC", num(m50["pr_auc"])),
        ("Recall", pct(m50["recall"])),
        ("Precision", pct(m50["precision"])),
        ("F1", num(m50["f1"])),
    ]
    for i, (label, value) in enumerate(kpis):
        left = Inches(0.7 + i * 3.1)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.4), Inches(2.9), Inches(1.6))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(0xEE, 0xF5, 0xFB)
        card.line.color.rgb = NAVY
        tf = card.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        set_run(r, size=14, color=GRAY)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = value
        set_run(r2, size=28, bold=True, color=NAVY)
    try_picture(s, "06_pr_curve_aleatorio.png", Inches(3.5), Inches(3.3), Inches(6))
    footer(s, 6)

    # 7 Top-K negocio
    s = blank_slide(prs)
    banner(s, "Métrica de negocio: ¿cuántos fraudes si revisamos Top-K?")
    try_picture(s, "08_captura_topk_aleatorio.png", Inches(0.5), Inches(1.2), Inches(7))
    box = s.shapes.add_textbox(Inches(7.7), Inches(1.8), Inches(5), Inches(4.5))
    add_bullets(box, [
        f"Top-100 → {cap[100]['frauds_found']} fraudes",
        f"Precision@100 = {pct(cap[100]['precision_at_k'])}",
        f"Top-500 → {cap[500]['frauds_found']} fraudes",
        f"Lift@500 = {num(cap[500]['lift'], 1)}×",
        "Traduce el score a una cola de revisión realista para el equipo.",
    ], size=16)
    footer(s, 7)

    # 8 Carga de alertas
    s = blank_slide(prs)
    banner(s, "No alcanza con FPR: mirar Precision y carga de revisión")
    try_picture(s, "09_workload_aleatorio.png", Inches(0.4), Inches(1.2), Inches(7.2))
    box = s.shapes.add_textbox(Inches(7.8), Inches(1.6), Inches(5), Inches(5))
    add_bullets(box, [
        f"FPR 5% sobre ~{m50['n_neg']:,} legítimas = miles de falsas alertas.",
        f"Umbral elegido: {msel['threshold']:.2f}",
        f"Alertas a revisar: {msel['alerts']}",
        f"Recall {pct(msel['recall'])} · Precision {pct(msel['precision'])}",
        f"TP={msel['tp']} · FP={msel['fp']} · FN={msel['fn']}",
    ], size=15)
    footer(s, 8)

    # 9 Temporal vs aleatorio
    s = blank_slide(prs)
    banner(s, "Validación temporal vs aleatoria")
    try_picture(s, "11_comparison_pr_auc.png", Inches(0.5), Inches(1.2), Inches(7.5))
    box = s.shapes.add_textbox(Inches(8.2), Inches(1.8), Inches(4.6), Inches(4.5))
    add_bullets(box, [
        "Aleatorio: mezcla pasado/futuro.",
        "Temporal: entrena con anteriores, prueba con posteriores.",
        f"Mejor temporal: {best_t}",
        f"PR-AUC temporal: {num(m50_t['pr_auc'])}",
        "Más realista para despliegue.",
    ], size=15)
    footer(s, 9)

    # 10 Impacto
    s = blank_slide(prs)
    banner(s, "Impacto esperado")
    box = s.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.5), Inches(5.5))
    add_bullets(box, [
        "Prioriza la cola de revisión humana (Top-100 / Top-500).",
        "Reduce fraudes no detectados vs reglas estáticas.",
        "Evita saturar al equipo con miles de falsas alarmas.",
        "Umbral = decisión de negocio explícita (Precision + capacidad).",
        "Escalable a scoring batch / near real-time.",
    ], size=17)
    footer(s, 10)

    # 11 Riesgos y ética
    s = blank_slide(prs)
    banner(s, "Riesgos, limitaciones y ética")
    box = s.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.5), Inches(5.5))
    add_bullets(box, [
        "Dataset de solo 2 días (2013): representatividad limitada.",
        "Variables PCA dificultan explicación regulatoria.",
        "Falso positivo = fricción al cliente; falso negativo = pérdida.",
        "Datos anonimizados; sin PII; licencia ODbL.",
        "Requiere monitoreo de drift y recalibración del umbral.",
    ], size=17)
    footer(s, 11)

    # 12 Cierre
    s = blank_slide(prs)
    banner(s, "Conclusiones y próximos pasos")
    box = s.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.5), Inches(5.5))
    add_bullets(box, [
        f"Modelo recomendado: {best_name} (mejor PR-AUC).",
        "KPI técnicos: PR-AUC · Recall · Precision.",
        "KPI de negocio: captura Top-100 / Top-500 y alertas revisables.",
        "Validar siempre con esquema temporal antes de producción.",
        "Próximos pasos: feedback del analista, drift monitoring, features adicionales.",
        "Código + informe + esta presentación en el repositorio del Grupo 33.",
    ], size=16)
    footer(s, 12)

    prs.save(str(OUT_PPTX))
    print(f"PPTX generado: {OUT_PPTX}")


def export_pdf():
    """Intenta exportar PDF vía LibreOffice; si no, genera PDF simple con reportlab + screenshots de slides no disponibles.
    Fallback: PDF con las figuras clave y bullets (misma narrativa).
    """
    import shutil
    import subprocess

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(OUT_PDF.parent), str(OUT_PPTX)],
                check=True,
                capture_output=True,
                timeout=120,
            )
            # LibreOffice usa el mismo stem
            generated = OUT_PPTX.with_suffix(".pdf")
            if generated.exists() and generated != OUT_PDF:
                generated.replace(OUT_PDF)
            if OUT_PDF.exists():
                print(f"PDF generado vía LibreOffice: {OUT_PDF}")
                return
        except Exception as exc:
            print(f"LibreOffice falló ({exc}); usando fallback reportlab.")

    # Fallback: PDF ejecutivo con reportlab
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib import colors

    data = json.loads(RESULTS.read_text(encoding="utf-8"))
    eda = data["eda"]
    rand = data["random_split"]
    best_name = rand["best_model"]
    best = rand["models"][best_name]
    m50 = best["metrics_thr_0.5"]
    msel = best["metrics_selected"]
    cap = {r["k"]: r for r in best["captura_at_k"]}
    temp = data["temporal_split"]
    best_t = temp["best_model"]
    m50_t = temp["models"][best_t]["metrics_thr_0.5"]

    title_s = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=22, textColor=colors.HexColor("#1F4E79"), alignment=TA_CENTER, spaceAfter=12)
    h_s = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=16, textColor=colors.HexColor("#1F4E79"), spaceAfter=8)
    b_s = ParagraphStyle("b", fontName="Helvetica", fontSize=12, leading=16, spaceAfter=4)

    story = []
    story.append(Spacer(1, 2 * cm))
    story.append(Paragraph("PROYECTO FINAL INTEGRADOR — GRUPO 33", title_s))
    story.append(Paragraph("Detección de Fraude en Tarjetas de Crédito", title_s))
    story.append(Paragraph("Mariana Chaile · Eugenia Salcedo Aisa · Sol Kidonakis · Marco Garcia · Esteban Derghazarian", b_s))
    story.append(PageBreak())

    slides_content = [
        ("Problema", [
            f"{eda['n_rows']:,} transacciones · {eda['n_fraud']} fraudes ({pct(eda['fraud_rate'], 3)}).",
            "Objetivo: maximizar detección sin saturar al equipo.",
            "Métricas: PR-AUC, Recall, Precision + captura Top-K.",
        ], None),
        ("Datos", [
            "Desbalance extremo; sin nulos; Amount_log; Time para split temporal.",
        ], "01_class_distribution.png"),
        ("Resultados", [
            f"Modelo: {best_name}",
            f"PR-AUC={num(m50['pr_auc'])} · Recall={pct(m50['recall'])} · Precision={pct(m50['precision'])}",
        ], "06_pr_curve_aleatorio.png"),
        ("Top-K", [
            f"Top-100 → {cap[100]['frauds_found']} fraudes",
            f"Top-500 → {cap[500]['frauds_found']} fraudes (lift {num(cap[500]['lift'], 1)}×)",
        ], "08_captura_topk_aleatorio.png"),
        ("Carga de alertas", [
            f"Umbral {msel['threshold']:.2f} → {msel['alerts']} alertas",
            f"Recall {pct(msel['recall'])} · Precision {pct(msel['precision'])}",
            "No guiarse solo por FPR.",
        ], "09_workload_aleatorio.png"),
        ("Temporal vs aleatorio", [
            f"Mejor temporal: {best_t} · PR-AUC={num(m50_t['pr_auc'])}",
            "Más realista para despliegue.",
        ], "11_comparison_pr_auc.png"),
        ("Conclusiones", [
            f"Recomendamos {best_name}.",
            "KPI negocio: Top-100/500 + alertas revisables.",
            "Validar siempre con esquema temporal.",
            "Próximos: feedback analista, drift, features adicionales.",
        ], None),
    ]

    for title, bullets, fig in slides_content:
        story.append(Paragraph(title, h_s))
        for b in bullets:
            story.append(Paragraph(f"• {b}", b_s))
        if fig and (FIGURES / fig).exists():
            img = Image(str(FIGURES / fig), width=16 * cm, height=8 * cm)
            # keep aspect
            from PIL import Image as PILImage
            with PILImage.open(FIGURES / fig) as im:
                w, h = im.size
            img.drawWidth = 16 * cm
            img.drawHeight = 16 * cm * h / w
            if img.drawHeight > 10 * cm:
                s = 10 * cm / img.drawHeight
                img.drawWidth *= s
                img.drawHeight *= s
            story.append(Spacer(1, 0.3 * cm))
            story.append(img)
        story.append(PageBreak())

    doc = SimpleDocTemplate(
        str(OUT_PDF),
        pagesize=landscape(A4),
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        title="Presentación Final — Grupo 33",
    )
    doc.build(story)
    print(f"PDF fallback generado: {OUT_PDF}")


def main():
    data = json.loads(RESULTS.read_text(encoding="utf-8"))
    build_pptx(data)
    export_pdf()


if __name__ == "__main__":
    main()
