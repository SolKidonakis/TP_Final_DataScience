"""Presentación Final del Proyecto (único PDF de entrega, máx. 12 páginas).

Genera PPTX y lo convierte a PDF vía LibreOffice.
Incluye índice, storytelling con cifras clave e imágenes con proporción correcta.
Sin apéndices ni código.
"""

from __future__ import annotations

import json
import shutil
import subprocess
from pathlib import Path

from PIL import Image as PILImage
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "reports" / "results.json"
FIGURES = ROOT / "reports" / "figures"
OUT_PPTX = ROOT / "reports" / "Presentacion_Final_Grupo33.pptx"
OUT_PDF = ROOT / "reports" / "Presentacion_Final_Grupo33.pdf"
# Copia también como "informe" unificado para el usuario
OUT_PDF_ALIAS = ROOT / "reports" / "Informe_Final_Grupo33.pdf"

NAVY = RGBColor(0x1F, 0x4E, 0x79)
TEAL = RGBColor(0x2E, 0x75, 0xB6)
DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xEE, 0xF5, 0xFB)
ACCENT = RGBColor(0xC0, 0x39, 0x2B)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TOTAL = 12


def pct(x, d=1):
    return f"{100 * float(x):.{d}f}%"


def num(x, d=3):
    return f"{float(x):.{d}f}"


def run(p, text, size=16, bold=False, color=DARK, font="Calibri"):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    return r


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def banner(slide, title: str):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.85))
    shape.fill.solid()
    shape.fill.fore_color.rgb = NAVY
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.35)
    tf.margin_top = Inches(0.18)
    p = tf.paragraphs[0]
    run(p, title, size=22, bold=True, color=WHITE)


def footer(slide, page: int):
    box = slide.shapes.add_textbox(Inches(0.4), Inches(7.15), Inches(12.5), Inches(0.28))
    p = box.text_frame.paragraphs[0]
    run(p, f"Grupo 33 · Detección de Fraude en Tarjetas de Crédito · {page}/{TOTAL}", size=10, color=GRAY)


def bullets(slide, left, top, width, height, items, size=16):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        run(p, f"•  {item}", size=size, color=DARK)
    return box


def add_picture_fit(slide, name: str, left, top, max_width, max_height):
    """Inserta imagen manteniendo aspecto, cabiendo en el rectángulo dado."""
    path = FIGURES / name
    if not path.exists():
        return False
    with PILImage.open(path) as im:
        pw, ph = im.size
    aspect = ph / pw
    width = max_width
    height = width * aspect
    if height > max_height:
        height = max_height
        width = height / aspect
    slide.shapes.add_picture(str(path), left, top, width=int(width), height=int(height))
    return True


def kpi_card(slide, left, top, w, h, label, value):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = LIGHT
    card.line.color.rgb = NAVY
    tf = card.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run(tf.paragraphs[0], label, size=12, color=GRAY)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    run(p2, value, size=24, bold=True, color=NAVY)


def build(data):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    eda = data["eda"]
    rand = data["random_split"]
    temp = data["temporal_split"]
    best_name = rand["best_model"]
    best = rand["models"][best_name]
    m50 = best["metrics_thr_0.5"]
    msel = best["metrics_selected"]
    cap = {r["k"]: r for r in best["captura_at_k"]}
    best_t = temp["best_model"]
    m50_t = temp["models"][best_t]["metrics_thr_0.5"]
    cap_t = {r["k"]: r for r in temp["models"][best_t]["captura_at_k"]}

    # ========== 1. Portada ==========
    s = blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    box = s.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(4.8))
    tf = box.text_frame
    tf.word_wrap = True
    lines = [
        ("DIPLOMADO EN CIENCIA DE DATOS Y ANÁLISIS AVANZADO", 14, False),
        ("Unidad N° 9 — Proyecto Final Integrador", 14, False),
        ("", 10, False),
        ("Detección de Fraude\nen Tarjetas de Crédito", 34, True),
        ("", 10, False),
        ("Presentación Final del Proyecto — Grupo 33", 18, True),
        ("Mariana Chaile · Eugenia Salcedo Aisa · Sol Kidonakis", 14, False),
        ("Marco Garcia · Esteban Derghazarian", 14, False),
        ("Octubre 2026", 14, False),
    ]
    first = True
    for text, size, bold in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(4)
        if text:
            run(p, text, size=size, bold=bold, color=WHITE)

    # ========== 2. Índice ==========
    s = blank(prs)
    banner(s, "Índice")
    bullets(
        s,
        Inches(1.2),
        Inches(1.3),
        Inches(10.5),
        Inches(5.5),
        [
            "1. Resumen ejecutivo",
            "2. Problema de negocio y objetivo",
            "3. Datos y metodología (CRISP-DM)",
            "4. Modelos evaluados y selección",
            "5. Resultados clave (PR-AUC, Recall, Precision)",
            "6. Métrica de negocio: captura Top-100 / Top-500",
            "7. Umbral operativo y carga de revisión",
            "8. Validación temporal vs aleatoria",
            "9. Impacto, riesgos y ética",
            "10. Conclusiones y recomendaciones",
        ],
        size=18,
    )
    footer(s, 2)

    # ========== 3. Resumen ejecutivo ==========
    s = blank(prs)
    banner(s, "1. Resumen ejecutivo")
    bullets(
        s,
        Inches(0.6),
        Inches(1.15),
        Inches(12.1),
        Inches(5.6),
        [
            f"Problema: detectar fraude en un flujo con desbalance extremo "
            f"(492 fraudes en 284.807 transacciones; 0,172 %).",
            f"Solución: modelo supervisado {best_name} que prioriza una cola de revisión humana.",
            f"Resultados clave (holdout aleatorio): PR-AUC {num(m50['pr_auc'])} · "
            f"Recall {pct(m50['recall'])} · Precision {pct(m50['precision'])}.",
            f"Negocio: Top-100 → {cap[100]['frauds_found']} fraudes; "
            f"Top-500 → {cap[500]['frauds_found']} fraudes (lift {num(cap[500]['lift'], 1)}×).",
            f"Operación: umbral {msel['threshold']:.2f} → solo {msel['alerts']} alertas "
            f"({msel['tp']} TP / {msel['fp']} FP), frente a miles con un FPR del 5 %.",
            f"Robustez: en validación temporal, {best_t} alcanza PR-AUC {num(m50_t['pr_auc'])}.",
        ],
        size=17,
    )
    footer(s, 3)

    # ========== 4. Problema ==========
    s = blank(prs)
    banner(s, "2. Problema de negocio y objetivo")
    bullets(
        s,
        Inches(0.6),
        Inches(1.2),
        Inches(12.1),
        Inches(5.5),
        [
            "Audiencia: Prevención de Fraude / Gestión de Riesgo Crediticio.",
            "Objetivo medible: clasificar cada transacción (fraude vs legítima) "
            "maximizando detección y controlando alertas revisables.",
            "Hipótesis: patrones en monto, tiempo y componentes PCA superan reglas estáticas.",
            "KPI técnicos: PR-AUC, Recall y Precision (no Accuracy ni solo FPR).",
            "KPI de negocio: ¿cuántos fraudes encontramos si revisamos Top-100 o Top-500?",
            "Restricción crítica: con tan pocos fraudes, FPR ≤ 5 % generaría miles de falsas alertas.",
        ],
        size=17,
    )
    footer(s, 4)

    # ========== 5. Datos y metodología ==========
    s = blank(prs)
    banner(s, "3. Datos y metodología (CRISP-DM)")
    add_picture_fit(s, "01_class_distribution.png", Inches(0.45), Inches(1.15), Inches(6.0), Inches(4.8))
    bullets(
        s,
        Inches(6.7),
        Inches(1.2),
        Inches(6.1),
        Inches(5.5),
        [
            "Fuente: ULB / Kaggle (ODbL).",
            f"Tras limpieza: {eda['n_rows']:,} filas; {eda['n_fraud']} fraudes.",
            "Sin nulos; duplicados eliminados.",
            "Amount_log = log1p(Amount).",
            "Scaler de Time y Amount_log dentro del Pipeline (sin leakage).",
            "Validación: (A) aleatorio 80/20 estratificado; (B) temporal por Time.",
        ],
        size=15,
    )
    footer(s, 5)

    # ========== 6. Modelos ==========
    s = blank(prs)
    banner(s, "4. Modelos evaluados y criterio de selección")
    bullets(
        s,
        Inches(0.6),
        Inches(1.15),
        Inches(12.1),
        Inches(2.4),
        [
            "Baseline: Regresión Logística (class_weight=balanced).",
            "Alternativas: Random Forest · XGBoost · LightGBM · XGBoost+SMOTE.",
            "Selección por PR-AUC (ranking bajo desbalance), con Recall/Precision y carga de alertas.",
            f"Modelo elegido: {best_name} — mejor equilibrio detección / alertas útiles.",
        ],
        size=16,
    )
    # mini tabla como bullets de cifras
    rows = []
    for name, m in rand["models"].items():
        m0 = m["metrics_thr_0.5"]
        mark = " ← elegido" if name == best_name else ""
        rows.append(
            f"{name}: PR-AUC {num(m0['pr_auc'])} · R {pct(m0['recall'])} · "
            f"P {pct(m0['precision'])}{mark}"
        )
    bullets(s, Inches(0.6), Inches(3.7), Inches(12.1), Inches(3.2), rows, size=14)
    footer(s, 6)

    # ========== 7. Resultados clave ==========
    s = blank(prs)
    banner(s, "5. Resultados clave — holdout aleatorio")
    w = Inches(2.85)
    for i, (lab, val) in enumerate([
        ("PR-AUC", num(m50["pr_auc"])),
        ("Recall", pct(m50["recall"])),
        ("Precision", pct(m50["precision"])),
        ("F1", num(m50["f1"])),
    ]):
        kpi_card(s, Inches(0.55 + i * 3.15), Inches(1.15), w, Inches(1.35), lab, val)
    add_picture_fit(s, "06_pr_curve_aleatorio.png", Inches(0.6), Inches(2.75), Inches(5.8), Inches(4.0))
    add_picture_fit(s, "07_confusion_aleatorio.png", Inches(7.0), Inches(2.75), Inches(5.5), Inches(4.0))
    footer(s, 7)

    # ========== 8. Top-K ==========
    s = blank(prs)
    banner(s, "6. Métrica de negocio: si revisamos Top-K, ¿cuántos fraudes?")
    add_picture_fit(s, "08_captura_topk_aleatorio.png", Inches(0.4), Inches(1.15), Inches(7.2), Inches(5.5))
    bullets(
        s,
        Inches(7.8),
        Inches(1.5),
        Inches(5.0),
        Inches(5.0),
        [
            f"Top-100 → {cap[100]['frauds_found']} fraudes",
            f"Precision@100 = {pct(cap[100]['precision_at_k'])}",
            f"Recall@100 = {pct(cap[100]['recall_at_k'])}",
            f"Top-500 → {cap[500]['frauds_found']} fraudes",
            f"Lift@500 = {num(cap[500]['lift'], 1)}×",
            "El score ordena la cola de revisión del equipo.",
        ],
        size=16,
    )
    footer(s, 8)

    # ========== 9. Umbral / carga ==========
    s = blank(prs)
    banner(s, "7. Umbral operativo: Precision y carga real de alertas")
    add_picture_fit(s, "09_workload_aleatorio.png", Inches(0.4), Inches(1.15), Inches(7.2), Inches(5.5))
    bullets(
        s,
        Inches(7.8),
        Inches(1.4),
        Inches(5.0),
        Inches(5.2),
        [
            f"FPR 5 % ≈ {int(0.05 * m50['n_neg']):,} alertas falsas solo en el holdout.",
            f"Umbral elegido: {msel['threshold']:.2f}",
            f"Alertas a revisar: {msel['alerts']}",
            f"Recall {pct(msel['recall'])} · Precision {pct(msel['precision'])}",
            f"TP={msel['tp']} · FP={msel['fp']} · FN={msel['fn']}",
            "No alcanza con controlar FPR: hay que mirar Precision y volumen.",
        ],
        size=15,
    )
    footer(s, 9)

    # ========== 10. Temporal ==========
    s = blank(prs)
    banner(s, "8. Validación temporal vs aleatoria")
    add_picture_fit(s, "11_comparison_pr_auc.png", Inches(0.35), Inches(1.15), Inches(7.5), Inches(5.5))
    bullets(
        s,
        Inches(8.0),
        Inches(1.4),
        Inches(4.8),
        Inches(5.2),
        [
            "Aleatorio: mezcla pasado y futuro.",
            "Temporal: entrena con anteriores, prueba con posteriores (más realista).",
            f"Mejor temporal: {best_t}",
            f"PR-AUC {num(m50_t['pr_auc'])}",
            f"Recall {pct(m50_t['recall'])} · Precision {pct(m50_t['precision'])}",
            f"Top-100 → {cap_t[100]['frauds_found']} · Top-500 → {cap_t[500]['frauds_found']}",
        ],
        size=15,
    )
    footer(s, 10)

    # ========== 11. Impacto / ética ==========
    s = blank(prs)
    banner(s, "9. Impacto, riesgos y consideraciones éticas")
    bullets(
        s,
        Inches(0.6),
        Inches(1.2),
        Inches(12.1),
        Inches(5.6),
        [
            "Impacto: prioriza la cola de revisión; reduce fraudes no detectados y el ruido operativo.",
            "Combinable con reglas existentes (ML como capa de priorización).",
            "Riesgos: dataset de 2 días (2013); PCA poco interpretable; posible drift.",
            "Ética: datos anonimizados (sin PII); falso positivo = fricción al cliente; "
            "falso negativo = pérdida financiera.",
            "El umbral es una decisión de negocio explícita, no un default técnico.",
            "Uso académico bajo licencia ODbL; despliegue real requiere supervisión humana.",
        ],
        size=16,
    )
    footer(s, 11)

    # ========== 12. Conclusiones ==========
    s = blank(prs)
    banner(s, "10. Conclusiones y recomendaciones")
    bullets(
        s,
        Inches(0.6),
        Inches(1.2),
        Inches(12.1),
        Inches(5.6),
        [
            f"Recomendamos {best_name} como modelo candidato (mejor PR-AUC en holdout aleatorio).",
            "Reportar siempre PR-AUC, Recall, Precision y cantidad de alertas.",
            "Usar captura Top-100 / Top-500 como KPI operativo de la cola de revisión.",
            "Validar con esquema temporal antes de cualquier despliegue.",
            "Próximos pasos: feedback del analista, monitoreo de drift, features adicionales "
            "(dispositivo, geo, historial) y calibración de scores.",
            "Código reproducible en el repositorio GitHub del Grupo 33 (README + notebook).",
        ],
        size=16,
    )
    footer(s, 12)

    prs.save(str(OUT_PPTX))
    print(f"PPTX: {OUT_PPTX}")


def export_pdf():
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice no disponible para convertir PPTX→PDF")
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(OUT_PDF.parent), str(OUT_PPTX)],
        check=True,
        capture_output=True,
        timeout=180,
    )
    generated = OUT_PPTX.with_suffix(".pdf")
    if generated.exists():
        # asegurar nombre canónico
        if generated.resolve() != OUT_PDF.resolve():
            generated.replace(OUT_PDF)
    # Alias: el "informe" de entrega es esta misma presentación
    shutil.copy2(OUT_PDF, OUT_PDF_ALIAS)
    # Copias en carpeta padre del TP
    parent = ROOT.parent
    shutil.copy2(OUT_PDF, parent / "Presentacion_Final_Grupo33.pdf")
    shutil.copy2(OUT_PDF_ALIAS, parent / "Informe_Final_Grupo33.pdf")
    print(f"PDF: {OUT_PDF}")
    print(f"Alias informe: {OUT_PDF_ALIAS}")


def main():
    data = json.loads(RESULTS.read_text(encoding="utf-8"))
    build(data)
    export_pdf()
    # Verificación de páginas
    try:
        from pypdf import PdfReader

        n = len(PdfReader(str(OUT_PDF)).pages)
        print(f"Páginas PDF: {n} (máx. permitido: 12)")
        if n > 12:
            raise SystemExit("ERROR: el PDF supera las 12 páginas")
    except ImportError:
        pass


if __name__ == "__main__":
    main()
